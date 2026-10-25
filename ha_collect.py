#!/usr/bin/env python3
"""Collect values of Home Assistant entities as Excel, CSV, or JSON-lines files.

Setup:
  export HA_URL="http://localhost:8123"                   # or http://<host-ip>:8123
  export HA_TOKEN="your-long-lived-token"

Usage:
  python3 ha_collect.py [--config FILE]                   # Collect values from entities
  python3 ha_collect.py [--config FILE] --list > info.md  # List entity IDs
"""

import argparse
import json
import logging
import os
import sys
import time
import urllib.error
import urllib.request
from pathlib import Path
from typing import Literal
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

import pandas as pd
import yaml
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, ConfigDict, Field, field_validator, model_validator

SECONDS_PER_DAY = 24 * 3600
TIMESTAMP_COLUMN = "__timestamp__"

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s %(levelname)-7s %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S",
)
log = logging.getLogger(__name__)


class Entity(BaseModel):
    model_config = ConfigDict(extra="forbid")
    id: str
    column: str = Field(min_length=1)
    unit: str | None


class Config(BaseModel):
    model_config = ConfigDict(extra="forbid")
    timezone: str
    samples_per_record: int = Field(gt=0)
    records_per_day: int = Field(gt=0, le=SECONDS_PER_DAY)
    timestamp_column: str = Field(min_length=1, default="timestamp")
    output_dir: Path
    output_format: Literal["excel", "csv", "jsonl"]
    entities: list[Entity] = Field(min_length=1)

    @property
    def sampling_interval(self):
        return SECONDS_PER_DAY / self.records_per_day / self.samples_per_record

    @field_validator("timezone")
    @classmethod
    def valid_timezone(cls, v):
        try:
            ZoneInfo(v)
        except ZoneInfoNotFoundError:
            raise ValueError(f"Unknown timezone: {v!r}")
        return v

    @field_validator("output_dir")
    @classmethod
    def dir_exists(cls, v):
        if not v.is_dir():
            raise ValueError(f"Directory does not exist: {v}")
        return v

    @model_validator(mode="after")
    def cross_checks(self):
        ids = [e.id for e in self.entities]
        columns = [e.column + e.unit for e in self.entities] + [self.timestamp_column]
        if len(set(ids)) != len(ids):
            raise ValueError("Duplicate entity ids")
        if len(set(columns)) != len(columns):
            raise ValueError("Duplicate column names")
        if TIMESTAMP_COLUMN in columns:
            raise ValueError(f"Disallowed column name: {TIMESTAMP_COLUMN}")
        if self.sampling_interval < 1:
            raise ValueError("Effective sampling interval is < 1s")
        return self


class HomeAssistantAPI:
    def __init__(self, url, token):
        self.__url = url
        self.__token = token

    def __get(self, path):
        req = urllib.request.Request(
            f"{self.__url}/api/{path}",
            headers={"Authorization": f"Bearer {self.__token}"},
            method="GET",
        )
        try:
            with urllib.request.urlopen(req, timeout=10) as resp:
                return json.load(resp)
        except urllib.error.HTTPError as e:
            if e.code == 401:
                log.error("401 Unauthorized: check token")
                sys.exit(os.EX_CONFIG)
            if e.code == 404:
                log.error("404 Not Found: check URL")
                sys.exit(os.EX_CONFIG)
            log.error("HTTP error %d from %s", e.code, req.full_url)
        except urllib.error.URLError as e:
            log.error("Cannot reach %s: %s", self.__url, e.reason)
        return None

    def get_entity_values(self):
        states = self.__get("states") or []
        entities = {}
        for s in states:
            entity_id = s["entity_id"]
            value = s["state"]
            unit = s["attributes"].get("unit_of_measurement")
            entities[entity_id] = value, unit
        return entities


class HomeAssistantSampler:
    def __init__(
        self,
        ha_api,
        on_record,
        entities,
        interval,
        samples_per_record,
    ):
        self.__ha_api = ha_api
        self.__on_record = on_record
        self.__samples_per_record = samples_per_record
        self.__interval = interval
        self.__entities_by_id = {e.id: e for e in entities}

    def __get_full_entity_column(self, entity):
        unit = entity.unit
        column = entity.column
        if len(unit) > 0:
            column = f"{column} [{unit}]"
        return column

    def __get_sample(self):
        entity_values = self.__ha_api.get_entity_values()
        sample = {TIMESTAMP_COLUMN: time.time()} | {
            self.__get_full_entity_column(e): None
            for e in self.__entities_by_id.values()
        }
        for entity_id, (value, _) in entity_values.items():
            if entity_id in self.__entities_by_id:
                entity = self.__entities_by_id[entity_id]
                sample[self.__get_full_entity_column(entity)] = value
        return sample

    def run(self):
        samples = []
        while True:
            start = time.perf_counter()
            log.info(
                "Taking sample %d/%d", (len(samples) + 1), self.__samples_per_record
            )
            samples.append(self.__get_sample())
            if len(samples) >= self.__samples_per_record:
                log.info("Aggregating %d samples into one record", len(samples))
                record = self.__aggregate_samples(samples)
                self.__on_record(record)
                samples.clear()
            sleep = self.__interval - (time.perf_counter() - start)
            try:
                time.sleep(sleep)
            except KeyboardInterrupt:
                return

    def __aggregate_samples(self, samples):
        df = pd.DataFrame(samples)
        record = {}
        missing_values = {"unavailable", "unknown", "none", ""}
        for col in df.columns:
            s = df[col]
            # Drop missing values
            valid = s[s.notna() & ~s.astype(str).str.lower().isin(missing_values)]
            # No valid values
            if valid.empty:
                record[col] = ""
                continue
            numeric = pd.to_numeric(valid, errors="coerce")
            if numeric.notna().all():
                # Aggregate numeric values using mean
                record[col] = float(numeric.mean())
            else:
                # Aggregate categorical values using mode
                record[col] = valid.mode().iloc[0]
        return record


def print_entities_list(ha_api, config_entities):
    entity_values = ha_api.get_entity_values()
    entities_by_id = {e.id: e for e in config_entities}
    entities = []
    for entity_id, (value, unit) in entity_values.items():
        in_config = entity_id in entities_by_id
        config_unit = None if not in_config else entities_by_id[entity_id].unit
        unit_missmatch = in_config and unit != config_unit
        entities.append(
            {
                "Entity Id": entity_id,
                "Value": value,
                "Unit": unit,
                "Selected In Config": "Yes" if in_config else "",
                "Config Error": (
                    f"Unit missmatch: {unit!r} vs. {config_unit!r}"
                    if unit_missmatch
                    else ""
                ),
            }
        )
    for e in config_entities:
        if e.id not in entity_values:
            entities.append(
                {
                    "Entity Id": e.id,
                    "Value": None,
                    "Unit": e.unit,
                    "Selected In Config": "Yes",
                    "Config Error": ("No such entity ID"),
                }
            )
    df = pd.DataFrame(entities).fillna("").sort_values("Entity Id")
    print(
        df.to_markdown(
            index=False,
            tablefmt="pipe",
            missingval="",
            disable_numparse=True,
        )
    )


def store_record_xlsx(output_dir, record, year, month):
    path = Path(output_dir) / f"{year:02}.xlsx"
    if path.exists():
        wb = load_workbook(path)
    else:
        wb = Workbook()
        wb.remove(wb.active)
    sheet = f"{year}-{month:02}"
    if sheet in wb.sheetnames:
        ws = wb[sheet]
    else:
        ws = wb.create_sheet(sheet)
        ws.append(list(record.keys()))
    ws.append(list(record.values()))
    ws.cell(ws.max_row, 1).number_format = "yyyy-mm-dd hh:mm:ss"
    wb.save(path)
    log.info("Appended record to %s", path)


def store_record_csv(output_dir, record, year, month):
    path = Path(output_dir) / f"{year}-{month:02}.csv"
    pd.DataFrame([record]).to_csv(
        path,
        mode="a",
        header=not path.exists(),
        index=False,
    )
    log.info("Appended record to %s", path)


def store_record_jsonl(output_dir, record, year, month):
    path = Path(output_dir) / f"{year}-{month:02}.jsonl"
    with open(
        path,
        "a",
        encoding="utf-8",
    ) as file:
        print(json.dumps(record, ensure_ascii=False), file=file)
    log.info("Appended record to %s", path)


def main():
    ha_url = os.environ.get("HA_URL", "http://localhost:8123").rstrip("/")
    ha_token = os.environ.get("HA_TOKEN", "")
    if not ha_token:
        sys.exit("Set the HA_TOKEN environment variable (long-lived access token).")

    p = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawTextHelpFormatter
    )
    p.add_argument(
        "-c",
        "--config",
        default="config.yaml",
        metavar="FILE",
        type=Path,
        help="Configuration YAML file",
    )
    p.add_argument(
        "-l",
        "--list",
        action="store_true",
        help="List available and selected entities",
    )
    args = p.parse_args()

    with open(args.config, "r") as file:
        config_dict = yaml.safe_load(file)
    config = Config(**config_dict)
    ha_api = HomeAssistantAPI(ha_url, ha_token)
    if args.list:
        print_entities_list(ha_api, config.entities)
        sys.exit()

    def store_record(record):
        record[TIMESTAMP_COLUMN] = (
            pd.to_datetime(record[TIMESTAMP_COLUMN], unit="s", utc=True)
            .tz_convert(config.timezone)
            .strftime("%Y-%m-%d %H:%M:%S")
        )
        year, month = record[TIMESTAMP_COLUMN].split("-")[:2]
        df = pd.DataFrame([record])
        df.columns = [config.timestamp_column] + list(df.columns[1:])
        record = df.round(3).iloc[0].to_dict()
        store_record_format = {
            "excel": store_record_xlsx,
            "csv": store_record_csv,
            "jsonl": store_record_jsonl,
        }.get(config.output_format)
        if store_record_format is None:
            raise ValueError("Unknown output format")
        store_record_format(config.output_dir, record, year, month)

    HomeAssistantSampler(
        ha_api,
        store_record,
        config.entities,
        config.sampling_interval,
        config.samples_per_record,
    ).run()


if __name__ == "__main__":
    main()
